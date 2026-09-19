"""Task 2 — parsing, indexing, safe editing, validation, dedupe. No LLM involved."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

from samples import make_samples
from task2 import workspace as wsmod
from task2.editor.docx_edit import EditError, apply_docx_ops
from task2.editor.generate import build_docx, build_xlsx
from task2.editor.validate import preflight_xlsx, validate, xlsx_stats
from task2.editor.xlsx_edit import apply_xlsx_ops
from task2.index import Index
from task2.parsers import parse_file
from task2.planner import normalise


@pytest.fixture(scope="module")
def corpus(tmp_path_factory):
    d = tmp_path_factory.mktemp("corpus")
    make_samples.contract_docx(d / "Vendor Contract - Nimbus Software.docx", "Nimbus Software", "INR 12,00,000", 30, "99.5%", 60, "the Nimbus CRM platform")
    make_samples.contract_docx(d / "Vendor Contract - Quartz Analytics.docx", "Quartz Analytics", "INR 7,20,000", 30, "99.9%", 60, "analytics platform")
    make_samples.invoice_register(d / "Invoice Register FY26.xlsx")
    make_samples.policy_pdf(d / "Vendor Policy.pdf")
    make_samples.board_minutes(d / "Board Minutes Q2 FY26.docx")
    (d / "notes.csv").write_text("Vendor,Owner\nNimbus Software,Garima\nQuartz Analytics,Ishan\n", encoding="utf-8")
    return d


# ----------------------------------------------------------------------------- parsers
def test_docx_units_and_outline(corpus):
    p = parse_file(corpus / "Vendor Contract - Nimbus Software.docx")
    assert p.kind == "docx" and p.outline[0].startswith("Master Services Agreement")
    assert any(u.kind == "table" for u in p.units) and any("within 30 days" in u.text for u in p.units)
    assert all(u.id.startswith(("p", "t")) for u in p.units)


def test_xlsx_hidden_sheet_and_formulas(corpus):
    p = parse_file(corpus / "Invoice Register FY26.xlsx")
    assert p.tables["Adjustments"]["hidden"] is True
    assert p.tables["Invoices"]["formulas"] > 100 and p.tables["Summary"]["formulas"] > 10
    assert "Terms" in p.tables["Invoices"]["columns"]
    assert any(u.kind == "rows" and "Net 45" in u.text for u in p.units), "row blocks must be searchable text"
    assert len(p.frames["Invoices"]) > 100


def test_pdf_pages_and_csv(corpus):
    p = parse_file(corpus / "Vendor Policy.pdf")
    assert p.kind == "pdf" and p.meta["pages"] >= 2 and any("99.9%" in u.text for u in p.units)
    c = parse_file(corpus / "notes.csv")
    assert c.kind == "csv" and c.tables["notes"]["rows"] == 2


# ----------------------------------------------------------------------------- index
def test_index_scoped_search(corpus, tmp_path):
    idx = Index(tmp_path)
    for i, name in enumerate(["Vendor Contract - Nimbus Software.docx", "Invoice Register FY26.xlsx", "Board Minutes Q2 FY26.docx"]):
        idx.add_file(f"f{i}", name, parse_file(corpus / name))
    assert idx.stats()["chunks"] > 10
    hits = idx.search("interest waived verbal agreement")
    assert hits and hits[0]["file_name"] == "Invoice Register FY26.xlsx" and hits[0]["loc"]["sheet"] == "Adjustments"
    scoped = idx.search("termination notice", file_ids=["f2"])
    assert scoped and all(h["file_id"] == "f2" for h in scoped)
    idx.remove_file("f1")
    assert not idx.search("Adjustments", file_ids=["f1"])


# ----------------------------------------------------------------------------- safe editing
def test_docx_edit_preserves_structure(corpus, tmp_path):
    src = corpus / "Vendor Contract - Nimbus Software.docx"
    parsed = parse_file(src)
    unit = next(u for u in parsed.units if "within 30 days" in u.text)
    dst = tmp_path / "edited.docx"
    res = apply_docx_ops(src, dst, [{"op": "replace_text", "unit": unit.id, "find": "within 30 days", "replace": "within 45 days"},
                                    {"op": "append_table_row", "unit": "t0", "values": ["Training", "Admin training", "2 days"]}])
    assert len(res["log"]) == 2
    report = validate(dst, original=src, expect=res["expect"])
    assert report["ok"], report
    again = parse_file(dst)
    edited = next(u for u in again.units if u.id == unit.id)
    assert "within 45 days" in edited.text and "within 30 days" not in edited.text
    assert again.outline == parsed.outline, "headings/styles must survive"
    with pytest.raises(EditError):
        apply_docx_ops(src, tmp_path / "x.docx", [{"op": "replace_text", "unit": unit.id, "find": "not in there", "replace": "y"}])


def test_xlsx_edit_refuses_formula_overwrite_and_keeps_everything(corpus, tmp_path):
    src = corpus / "Invoice Register FY26.xlsx"
    before = xlsx_stats(src)
    with pytest.raises(EditError):
        apply_xlsx_ops(src, tmp_path / "bad.xlsx", [{"op": "set_cell", "sheet": "Invoices", "cell": "J2", "value": 5}])
    dst = tmp_path / "edited.xlsx"
    res = apply_xlsx_ops(src, dst, [{"op": "update_rows", "sheet": "Invoices", "where": {"column": "Vendor", "equals": "Nimbus Software"}, "set": {"Terms": "Net 30"}},
                                    {"op": "add_sheet", "name": "Audit Notes", "rows": [["Finding", "Detail"], ["Terms", "Nimbus moved to Net 30"]]}], header_rows={"Invoices": 1})
    report = validate(dst, original=src, expect=res["expect"])
    assert report["ok"], report
    after = xlsx_stats(dst)
    assert after["formulas"] == before["formulas"] and "Adjustments" in after["hidden"] and "Audit Notes" in after["sheets"]
    wb = load_workbook(dst)
    terms = [wb["Invoices"].cell(r, 5).value for r in range(2, wb["Invoices"].max_row + 1) if wb["Invoices"].cell(r, 2).value == "Nimbus Software"]
    assert terms and all(t == "Net 30" for t in terms)


def test_preflight_blocks_workbooks_with_charts(tmp_path):
    wb = Workbook()
    ws = wb.active
    for r in [["a", 1], ["b", 2], ["c", 3]]:
        ws.append(r)
    ch = BarChart()
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=3))
    ws.add_chart(ch, "D2")
    p = tmp_path / "chart.xlsx"
    wb.save(p)
    assert preflight_xlsx(p), "workbooks with charts must not be edited in place"


def test_validate_catches_corruption(tmp_path):
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"PK\x03\x04 this is not a real docx")
    assert validate(bad)["ok"] is False


def test_generate_docx_and_xlsx(tmp_path):
    spec = {"title": "Memo", "sections": [{"heading": "Findings", "paragraphs": ["One."], "bullets": ["a", "b"], "table": {"columns": ["Item", "Value"], "rows": [["x", "1"]]}}]}
    build_docx(spec, tmp_path / "m.docx")
    assert validate(tmp_path / "m.docx")["ok"]
    build_xlsx({"sheets": [{"name": "Vendors", "columns": ["Vendor", "Fee"], "rows": [["A", 100], ["B", 250]], "total_row": True}]}, tmp_path / "v.xlsx")
    st = xlsx_stats(tmp_path / "v.xlsx")
    assert st["sheets"] == ["Vendors"] and st["formulas"] == 1


# ----------------------------------------------------------------------------- workspace
def test_workspace_dedupe_and_zip(tmp_path, monkeypatch):
    monkeypatch.setattr(wsmod, "WS_ROOT", tmp_path)
    ws = wsmod.Store().create("t")
    data = b"hello world"
    (r1, dup1), = ws.add_bytes("a.txt", data)
    (r2, dup2), = ws.add_bytes("copy of a.txt", data)
    assert dup1 is False and dup2 is True and r1.id == r2.id
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("folder/b.md", "# hi\n\ntext")
        z.writestr("__MACOSX/._b.md", "junk")
        z.writestr("folder/a.txt", data)
    out = ws.add_bytes("bundle.zip", buf.getvalue())
    assert len(out) == 2 and sum(1 for _, d in out if d) == 1  # a.txt again -> duplicate, b.md new
    v2 = ws.add_version(r1.id, b"hello world v2", "edited")
    assert v2.version == 2 and v2.parent_id == r1.id and r1 not in ws.latest_files() and v2 in ws.latest_files()


def test_planner_normalise_resolves_names(tmp_path, monkeypatch):
    monkeypatch.setattr(wsmod, "WS_ROOT", tmp_path)
    ws = wsmod.Store().create("t")
    for n in ("Vendor Contract - Nimbus Software.docx", "Board Minutes.docx", "Invoice Register.xlsx"):
        (rec, _), = ws.add_bytes(n, n.encode())
        rec.status = "indexed"
    p = normalise({"strategy": "cross_audit", "target_files": ["Vendor Contract - Nimbus Software.docx"], "support_files": "all", "queries": ["x"]}, ws, "audit the nimbus contract")
    assert p["strategy"] == "cross_audit" and len(p["target_files"]) == 1 and p["support_all"] and len(p["support_files"]) == 2
    p2 = normalise({"strategy": "nonsense", "target_files": [], "support_files": []}, ws, "what does the board minutes say")
    assert p2["strategy"] == "qa" and p2["target_files"] and ws.files[p2["target_files"][0]].name == "Board Minutes.docx"
