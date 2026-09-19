"""XLSX/XLSM/CSV -> every sheet (hidden ones included and flagged) becomes:
  * a DataFrame for SQL (DuckDB) — the model queries, never reads, big tables
  * a 'sheet' unit describing columns/size, and 'rows' units (blocks of 20 rows as a markdown table)
    so that clauses buried in a spreadsheet are still findable by search
Formulas are counted and sampled from a second, non-data_only load."""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import ParsedDoc, Unit

ROW_BLOCK = 20
MAX_ROWS_INDEXED = 2000


def _header_row(rows: list[list[Any]]) -> int:
    for i, r in enumerate(rows[:15]):
        if sum(1 for c in r if c not in (None, "")) >= 2:
            return i
    return 0


def _frame(rows: list[list[Any]]) -> tuple[pd.DataFrame, int]:
    if not rows:
        return pd.DataFrame(), 0
    h = _header_row(rows)
    header = [str(c).strip() if c not in (None, "") else f"col{j + 1}" for j, c in enumerate(rows[h])]
    seen: dict[str, int] = {}
    cols = []
    for c in header:
        seen[c] = seen.get(c, 0) + 1
        cols.append(c if seen[c] == 1 else f"{c}_{seen[c]}")
    body = [r + [None] * (len(cols) - len(r)) for r in rows[h + 1 :] if any(c not in (None, "") for c in r)]
    df = pd.DataFrame([r[: len(cols)] for r in body], columns=cols)
    return df, h + 1  # 1-based sheet row of the header


def _md(df: pd.DataFrame, start_idx: int, end_idx: int, header_row: int) -> str:
    sub = df.iloc[start_idx:end_idx]
    cols = list(sub.columns)[:14]
    lines = ["| row | " + " | ".join(str(c) for c in cols) + " |"]
    for i, (_, r) in enumerate(sub.iterrows()):
        sheet_row = header_row + 1 + start_idx + i
        lines.append(f"| {sheet_row} | " + " | ".join(_cell(r[c]) for c in cols) + " |")
    return "\n".join(lines)


def _cell(v: Any) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = str(v).replace("\n", " ").replace("|", "/")
    return s[:60]


def _sheet_units(name: str, df: pd.DataFrame, header_row: int, info: dict[str, Any]) -> list[Unit]:
    units = [Unit(id=f"{name}!sheet", kind="sheet",
                  text=f"Sheet '{name}'{' (HIDDEN)' if info.get('hidden') else ''}: {len(df)} rows × {len(df.columns)} columns. Columns: {', '.join(map(str, df.columns[:30]))}."
                       + (f" Contains {info['formulas']} formulas." if info.get("formulas") else ""),
                  loc={"sheet": name})]
    n = min(len(df), MAX_ROWS_INDEXED)
    for start in range(0, n, ROW_BLOCK):
        end = min(start + ROW_BLOCK, n)
        r1, r2 = header_row + 1 + start, header_row + end
        units.append(Unit(id=f"{name}!r{r1}-{r2}", kind="rows", text=f"Sheet '{name}' rows {r1}-{r2}:\n" + _md(df, start, end, header_row),
                          loc={"sheet": name, "rows": [r1, r2]}))
    return units


def parse_xlsx(path: Path) -> ParsedDoc:
    wb_vals = load_workbook(str(path), data_only=True, read_only=False)
    wb_form = load_workbook(str(path), data_only=False, read_only=False)
    units: list[Unit] = []
    outline: list[str] = []
    tables: dict[str, dict[str, Any]] = {}
    frames: dict[str, pd.DataFrame] = {}
    hidden: list[str] = []
    for ws in wb_vals.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        df, header_row = _frame(rows)
        wf = wb_form[ws.title]
        formulas = []
        for row in wf.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas.append(f"{c.coordinate}: {c.value}")
        is_hidden = ws.sheet_state != "visible"
        if is_hidden:
            hidden.append(ws.title)
        info = {
            "columns": [str(c) for c in df.columns], "rows": int(len(df)), "header_row": header_row, "hidden": is_hidden,
            "formulas": len(formulas), "formula_samples": formulas[:8],
            "sample": [{str(k): _cell(v) for k, v in r.items()} for _, r in df.head(3).iterrows()],
            "dims": f"{ws.max_row}x{ws.max_column}",
        }
        tables[ws.title] = info
        frames[ws.title] = df
        outline.append(f"{ws.title}{' (hidden)' if is_hidden else ''}: {', '.join(info['columns'][:8])}")
        units.extend(_sheet_units(ws.title, df, header_row, info))
    return ParsedDoc(kind="xlsx", units=units, outline=outline, tables=tables, frames=frames,
                     meta={"sheets": [ws.title for ws in wb_vals.worksheets], "hidden_sheets": hidden, "has_vba": path.suffix.lower() == ".xlsm"})


def parse_csv(path: Path) -> ParsedDoc:
    df = pd.read_csv(path)
    df = df.where(pd.notnull(df), None)
    name = path.stem
    info = {"columns": [str(c) for c in df.columns], "rows": int(len(df)), "header_row": 1, "hidden": False, "formulas": 0,
            "sample": [{str(k): _cell(v) for k, v in r.items()} for _, r in df.head(3).iterrows()]}
    return ParsedDoc(kind="csv", units=_sheet_units(name, df, 1, info), outline=[f"{name}: {', '.join(info['columns'][:8])}"],
                     tables={name: info}, frames={name: df}, meta={"sheets": [name], "hidden_sheets": []})
