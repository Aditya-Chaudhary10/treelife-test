"""DOCX -> paragraph/table units in document order (python-docx). Paragraph ids `p<i>` index into
`Document().paragraphs`, table ids `t<i>` into `Document().tables` — the same indices the editor uses."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from . import ParsedDoc, Unit


def iter_block_items(doc):
    """Yield ('p', index, Paragraph) / ('t', index, Table) in body order."""
    body = doc.element.body
    p_i = t_i = 0
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield "p", p_i, Paragraph(child, doc)
            p_i += 1
        elif tag == "tbl":
            yield "t", t_i, Table(child, doc)
            t_i += 1


def table_to_text(tbl: Table) -> str:
    rows = []
    for r in tbl.rows:
        cells = [c.text.strip().replace("\n", " ") for c in r.cells]
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def parse_docx(path: Path) -> ParsedDoc:
    doc = Document(str(path))
    units: list[Unit] = []
    outline: list[str] = []
    current_heading = ""
    for kind, idx, obj in iter_block_items(doc):
        if kind == "p":
            text = obj.text.strip()
            if not text:
                continue
            style = (obj.style.name if obj.style is not None else "") or ""
            is_heading = style.lower().startswith(("heading", "title"))
            if is_heading:
                current_heading = text
                if len(outline) < 60:
                    outline.append(text)
            units.append(Unit(id=f"p{idx}", kind="heading" if is_heading else "paragraph", text=text,
                              loc={"paragraph": idx, "style": style, "section": current_heading}))
        else:
            text = table_to_text(obj)
            if text.strip():
                units.append(Unit(id=f"t{idx}", kind="table", text=text, loc={"table": idx, "section": current_heading, "rows": len(obj.rows), "cols": len(obj.columns)}))
    meta = {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "sections": len(doc.sections)}
    core = doc.core_properties
    if core.title:
        meta["title"] = core.title
    return ParsedDoc(kind="docx", units=units, outline=outline, meta=meta)
