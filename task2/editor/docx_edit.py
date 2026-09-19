"""Apply structured edit operations to a DOCX with python-docx. The model never writes XML;
it names a paragraph/table by the ids produced at parse time and says what should change.

ops:
  {"op":"replace_text","unit":"p12","find":"30 days","replace":"45 days"}
  {"op":"set_paragraph_text","unit":"p12","text":"..."}            (keeps paragraph style + first run formatting)
  {"op":"insert_paragraph_after","unit":"p12","text":"...","style":"Normal"|null}
  {"op":"delete_paragraph","unit":"p12"}
  {"op":"set_table_cell","unit":"t0","row":1,"col":2,"text":"..."}
  {"op":"append_table_row","unit":"t0","values":["a","b","c"]}
  {"op":"replace_all","find":"...","replace":"..."}
"""
from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.text.paragraph import Paragraph


class EditError(ValueError):
    pass


def _unit_index(unit: str, prefix: str) -> int:
    m = re.fullmatch(rf"{prefix}(\d+)", str(unit).strip())
    if not m:
        raise EditError(f"unit '{unit}' is not a {prefix}-unit id")
    return int(m.group(1))


def _paragraph(doc, unit: str) -> Paragraph:
    i = _unit_index(unit, "p")
    if i >= len(doc.paragraphs):
        raise EditError(f"paragraph {unit} does not exist (document has {len(doc.paragraphs)})")
    return doc.paragraphs[i]


def _table(doc, unit: str):
    i = _unit_index(unit, "t")
    if i >= len(doc.tables):
        raise EditError(f"table {unit} does not exist (document has {len(doc.tables)})")
    return doc.tables[i]


def replace_in_paragraph(p: Paragraph, find: str, replace: str) -> int:
    """Replace inside runs when possible (keeps bold/italic per run); otherwise collapse runs
    into the first one, preserving the first run's formatting."""
    if find not in p.text:
        return 0
    n = 0
    for r in p.runs:
        if find in r.text:
            n += r.text.count(find)
            r.text = r.text.replace(find, replace)
    if n:
        return n
    # match spans several runs
    full = p.text
    n = full.count(find)
    new = full.replace(find, replace)
    if p.runs:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(new)
    return n


def set_paragraph_text(p: Paragraph, text: str) -> None:
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def insert_paragraph_after(p: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = copy.deepcopy(p._p)
    for child in list(new_p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            new_p.remove(child)
    p._p.addnext(new_p)
    para = Paragraph(new_p, p._parent)
    run = para.add_run(text)
    if p.runs:  # inherit first-run character formatting
        src = p.runs[0]
        run.bold, run.italic, run.underline = src.bold, src.italic, src.underline
        if src.font.size:
            run.font.size = src.font.size
        if src.font.name:
            run.font.name = src.font.name
    if style:
        try:
            para.style = style
        except Exception:
            pass
    return para


def apply_docx_ops(src: Path, dst: Path, ops: list[dict[str, Any]]) -> dict[str, Any]:
    doc = Document(str(src))
    log: list[str] = []
    paragraph_delta = 0
    for op in ops:
        kind = op.get("op")
        if kind == "replace_text":
            p = _paragraph(doc, op["unit"])
            n = replace_in_paragraph(p, str(op["find"]), str(op["replace"]))
            if n == 0:
                raise EditError(f"'{op['find']}' not found in {op['unit']}: \"{p.text[:80]}\"")
            log.append(f"{op['unit']}: replaced '{op['find']}' → '{op['replace']}' ({n}×)")
        elif kind == "set_paragraph_text":
            p = _paragraph(doc, op["unit"])
            old = p.text
            set_paragraph_text(p, str(op["text"]))
            log.append(f"{op['unit']}: rewrote paragraph (was: \"{old[:60]}…\")")
        elif kind == "insert_paragraph_after":
            p = _paragraph(doc, op["unit"])
            insert_paragraph_after(p, str(op["text"]), op.get("style"))
            paragraph_delta += 1
            log.append(f"{op['unit']}: inserted new paragraph after it")
        elif kind == "delete_paragraph":
            p = _paragraph(doc, op["unit"])
            old = p.text
            p._p.getparent().remove(p._p)
            paragraph_delta -= 1
            log.append(f"{op['unit']}: deleted paragraph \"{old[:60]}\"")
        elif kind == "set_table_cell":
            t = _table(doc, op["unit"])
            r, c = int(op["row"]), int(op["col"])
            if r >= len(t.rows) or c >= len(t.columns):
                raise EditError(f"cell ({r},{c}) outside table {op['unit']} ({len(t.rows)}x{len(t.columns)})")
            cell = t.cell(r, c)
            old = cell.text
            if cell.paragraphs:
                set_paragraph_text(cell.paragraphs[0], str(op["text"]))
                for extra in cell.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
            log.append(f"{op['unit']} cell({r},{c}): '{old}' → '{op['text']}'")
        elif kind == "append_table_row":
            t = _table(doc, op["unit"])
            vals = [str(v) for v in op.get("values", [])]
            cells = t.add_row().cells
            for i, v in enumerate(vals[: len(cells)]):
                cells[i].text = v
            log.append(f"{op['unit']}: appended row {vals}")
        elif kind == "replace_all":
            find, rep = str(op["find"]), str(op["replace"])
            n = 0
            for p in doc.paragraphs:
                n += replace_in_paragraph(p, find, rep)
            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            n += replace_in_paragraph(p, find, rep)
            if n == 0:
                raise EditError(f"'{find}' not found anywhere in the document")
            log.append(f"replaced '{find}' → '{rep}' in {n} place(s)")
        else:
            raise EditError(f"unknown docx op '{kind}'")
    doc.save(str(dst))
    return {"log": log, "expect": {"paragraph_delta": paragraph_delta}}
