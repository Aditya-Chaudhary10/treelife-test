"""Apply structured edit operations to a workbook with openpyxl (keep_vba for .xlsm).
Values only; formulas are never overwritten unless the op explicitly says so.

ops:
  {"op":"set_cell","sheet":"Invoices","cell":"D3","value":337500}
  {"op":"set_formula","sheet":"Summary","cell":"B9","formula":"=SUM(B2:B8)"}
  {"op":"append_row","sheet":"Invoices","values":[...]}
  {"op":"update_rows","sheet":"Invoices","where":{"column":"Vendor","equals":"Nimbus Software"},"set":{"Status":"Paid"}}
  {"op":"add_sheet","name":"Audit Notes","rows":[["Finding","Detail"],["...","..."]]}
  {"op":"unhide_sheet","sheet":"Adjustments"}
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .docx_edit import EditError


def _coerce(v: Any) -> Any:
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("="):
            return s
        try:
            if s.replace(",", "").replace(".", "", 1).lstrip("-").isdigit():
                return float(s.replace(",", "")) if "." in s else int(s.replace(",", ""))
        except ValueError:
            pass
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S"):
            try:
                return datetime.strptime(s, fmt)
            except ValueError:
                continue
    return v


def _sheet(wb, name: str):
    if name in wb.sheetnames:
        return wb[name]
    low = {s.lower(): s for s in wb.sheetnames}
    if name.lower() in low:
        return wb[low[name.lower()]]
    raise EditError(f"sheet '{name}' not found; sheets: {wb.sheetnames}")


def _header(ws, header_row: int) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[header_row] if c.value not in (None, "")}


def apply_xlsx_ops(src: Path, dst: Path, ops: list[dict[str, Any]], header_rows: dict[str, int] | None = None) -> dict[str, Any]:
    keep_vba = src.suffix.lower() == ".xlsm"
    wb = load_workbook(str(src), keep_vba=keep_vba)
    header_rows = header_rows or {}
    log: list[str] = []
    formulas_removed = formulas_added = 0
    for op in ops:
        kind = op.get("op")
        if kind == "set_cell":
            ws = _sheet(wb, op["sheet"])
            cell = ws[op["cell"]]
            old = cell.value
            if isinstance(old, str) and old.startswith("=") and not op.get("allow_formula_overwrite"):
                raise EditError(f"{op['sheet']}!{op['cell']} holds a formula ({old}); refusing to overwrite. Use set_formula or allow_formula_overwrite.")
            cell.value = _coerce(op.get("value"))
            log.append(f"{ws.title}!{op['cell']}: {old!r} → {cell.value!r}")
        elif kind == "set_formula":
            ws = _sheet(wb, op["sheet"])
            cell = ws[op["cell"]]
            old = cell.value
            f = str(op["formula"]).strip()
            if not f.startswith("="):
                f = "=" + f
            if not (isinstance(old, str) and old.startswith("=")):
                formulas_added += 1
            cell.value = f
            log.append(f"{ws.title}!{op['cell']}: formula set to {f} (was {old!r})")
        elif kind == "append_row":
            ws = _sheet(wb, op["sheet"])
            vals = [_coerce(v) for v in op.get("values", [])]
            ws.append(vals)
            log.append(f"{ws.title}: appended row {ws.max_row}: {vals}")
        elif kind == "update_rows":
            ws = _sheet(wb, op["sheet"])
            hr = int(header_rows.get(ws.title, 1))
            cols = _header(ws, hr)
            where, setv = op.get("where") or {}, op.get("set") or {}
            wc = cols.get(str(where.get("column")))
            if not wc:
                raise EditError(f"column '{where.get('column')}' not in {ws.title} header {list(cols)}")
            for k in setv:
                if k not in cols:
                    raise EditError(f"column '{k}' not in {ws.title} header {list(cols)}")
            target = str(where.get("equals", "")).strip().lower()
            n = 0
            for r in range(hr + 1, ws.max_row + 1):
                v = ws.cell(r, wc).value
                if v is not None and str(v).strip().lower() == target:
                    for k, nv in setv.items():
                        c = ws.cell(r, cols[k])
                        if isinstance(c.value, str) and c.value.startswith("="):
                            raise EditError(f"{ws.title}!{c.coordinate} holds a formula; refusing to overwrite")
                        c.value = _coerce(nv)
                    n += 1
            if n == 0:
                raise EditError(f"no rows where {where.get('column')} = '{where.get('equals')}' in {ws.title}")
            log.append(f"{ws.title}: updated {n} row(s) where {where.get('column')} = '{where.get('equals')}' → {setv}")
        elif kind == "add_sheet":
            name = str(op.get("name") or "Sheet")[:31]
            if name in wb.sheetnames:
                raise EditError(f"sheet '{name}' already exists")
            ws = wb.create_sheet(name)
            for row in op.get("rows") or []:
                ws.append([_coerce(v) for v in row])
            for i, _ in enumerate((op.get("rows") or [[]])[0], start=1):
                ws.column_dimensions[get_column_letter(i)].width = 28
            log.append(f"added sheet '{name}' with {len(op.get('rows') or [])} rows")
        elif kind == "unhide_sheet":
            ws = _sheet(wb, op["sheet"])
            ws.sheet_state = "visible"
            log.append(f"{ws.title}: made visible")
        else:
            raise EditError(f"unknown xlsx op '{kind}'")
    try:
        wb.calculation.fullCalcOnLoad = True  # Excel recomputes formulas when the file is opened
    except Exception:
        pass
    wb.save(str(dst))
    return {"log": log, "expect": {"formulas_removed": formulas_removed, "formulas_added": formulas_added}}
