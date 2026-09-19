"""File-safety checks. Every generated or edited file is re-opened with the real libraries and
compared against the original's invariants before it is offered for download."""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

RISKY_XLSX_PARTS = ("xl/charts/", "xl/drawings/", "xl/pivotTables/", "xl/pivotCache/", "xl/slicers/", "xl/embeddings/")


def zip_parts(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return z.namelist()


def zip_ok(path: Path) -> str | None:
    try:
        with zipfile.ZipFile(path) as z:
            bad = z.testzip()
            return f"corrupt member: {bad}" if bad else None
    except zipfile.BadZipFile as e:
        return f"not a valid OOXML zip: {e}"


def docx_stats(path: Path) -> dict[str, Any]:
    from docx import Document

    d = Document(str(path))
    parts = zip_parts(path)
    return {
        "paragraphs": len(d.paragraphs), "tables": len(d.tables), "sections": len(d.sections),
        "styles": len(d.styles), "has_document_xml": "word/document.xml" in parts, "has_styles": "word/styles.xml" in parts,
        "headers_footers": sum(1 for p in parts if p.startswith("word/header") or p.startswith("word/footer")),
        "images": sum(1 for p in parts if p.startswith("word/media/")),
    }


def xlsx_stats(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    parts = zip_parts(path)
    wb = load_workbook(str(path), data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    formulas = 0
    merged = 0
    for ws in wb.worksheets:
        merged += len(ws.merged_cells.ranges)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
    return {
        "sheets": wb.sheetnames, "hidden": [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"],
        "formulas": formulas, "merged_ranges": merged, "defined_names": len(list(wb.defined_names)) if hasattr(wb.defined_names, "__iter__") else 0,
        "has_vba": "xl/vbaProject.bin" in parts, "risky_parts": sorted({p.split("/")[1] for p in parts if p.startswith(RISKY_XLSX_PARTS)}),
    }


def preflight_xlsx(path: Path) -> list[str]:
    """Reasons NOT to auto-edit this workbook in place (openpyxl cannot round-trip these parts)."""
    st = xlsx_stats(path)
    return [f"workbook contains {p} which the in-place editor cannot preserve" for p in st["risky_parts"]]


def validate(path: Path, original: Path | None = None, expect: dict[str, Any] | None = None) -> dict[str, Any]:
    """Open the result, run integrity checks, compare with the original. Returns {ok, issues, warnings, stats}."""
    issues: list[str] = []
    warnings: list[str] = []
    err = zip_ok(path)
    if err:
        return {"ok": False, "issues": [err], "warnings": [], "stats": {}}
    ext = path.suffix.lower()
    try:
        stats = docx_stats(path) if ext == ".docx" else xlsx_stats(path)
    except Exception as e:
        return {"ok": False, "issues": [f"library cannot open the file: {type(e).__name__}: {e}"], "warnings": [], "stats": {}}
    expect = expect or {}
    if original is not None and original.exists():
        try:
            before = docx_stats(original) if ext == ".docx" else xlsx_stats(original)
        except Exception:
            before = {}
        if ext == ".docx" and before:
            if not stats["has_document_xml"] or not stats["has_styles"]:
                issues.append("document.xml or styles.xml missing")
            delta = expect.get("paragraph_delta", 0)
            if stats["paragraphs"] != before["paragraphs"] + delta:
                warnings.append(f"paragraph count changed {before['paragraphs']} → {stats['paragraphs']} (expected {before['paragraphs'] + delta})")
            if stats["tables"] < before["tables"]:
                issues.append(f"tables lost: {before['tables']} → {stats['tables']}")
            if stats["images"] < before["images"]:
                issues.append(f"images lost: {before['images']} → {stats['images']}")
            if stats["headers_footers"] < before["headers_footers"]:
                issues.append("headers/footers lost")
        if ext in (".xlsx", ".xlsm") and before:
            missing = [s for s in before["sheets"] if s not in stats["sheets"]]
            if missing:
                issues.append(f"sheets lost: {missing}")
            min_formulas = before["formulas"] - expect.get("formulas_removed", 0) + expect.get("formulas_added", 0)
            if stats["formulas"] < min_formulas:
                issues.append(f"formulas lost: {before['formulas']} → {stats['formulas']}")
            if before["has_vba"] and not stats["has_vba"]:
                issues.append("VBA project (macros) lost")
            if stats["merged_ranges"] < before["merged_ranges"]:
                issues.append("merged cell ranges lost")
            if before.get("risky_parts"):
                warnings.append(f"source had {before['risky_parts']} — verify they survived")
    return {"ok": not issues, "issues": issues, "warnings": warnings, "stats": stats}
