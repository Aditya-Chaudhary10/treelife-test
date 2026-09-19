"""Build NEW documents from a structured spec. The model produces content (headings, paragraphs,
bullets, tables); python-docx / openpyxl produce the bytes. No model-written XML ever.

docx spec: {"title": "...", "subtitle": "...", "sections": [{"heading": "...", "paragraphs": ["..."], "bullets": ["..."],
            "table": {"columns": ["..."], "rows": [["..."]]}}]}
xlsx spec: {"sheets": [{"name": "...", "columns": ["..."], "rows": [[...]], "total_row": true|false}]}
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .xlsx_edit import _coerce


def build_docx(spec: dict[str, Any], path: Path) -> None:
    d = Document()
    d.add_heading(str(spec.get("title") or "Document"), 0)
    if spec.get("subtitle"):
        p = d.add_paragraph(str(spec["subtitle"]))
        p.runs[0].italic = True
    for sec in spec.get("sections") or []:
        if sec.get("heading"):
            d.add_heading(str(sec["heading"]), 1)
        for para in sec.get("paragraphs") or []:
            d.add_paragraph(str(para))
        for b in sec.get("bullets") or []:
            d.add_paragraph(str(b), style="List Bullet")
        tbl = sec.get("table")
        if tbl and tbl.get("columns"):
            cols = [str(c) for c in tbl["columns"]]
            t = d.add_table(rows=1, cols=len(cols))
            t.style = "Table Grid"
            for i, c in enumerate(cols):
                cell = t.rows[0].cells[i]
                cell.text = c
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in tbl.get("rows") or []:
                cells = t.add_row().cells
                for i, v in enumerate(list(row)[: len(cols)]):
                    cells[i].text = "" if v is None else str(v)
            d.add_paragraph()
    style = d.styles["Normal"]
    style.font.size = Pt(11)
    d.save(str(path))


def build_xlsx(spec: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F766E")
    for sh in spec.get("sheets") or [{"name": "Sheet1", "columns": [], "rows": []}]:
        ws = wb.create_sheet(str(sh.get("name") or "Sheet")[:31])
        cols = [str(c) for c in sh.get("columns") or []]
        if cols:
            ws.append(cols)
            for c in ws[1]:
                c.font, c.fill = head_font, head_fill
            ws.freeze_panes = "A2"
        for row in sh.get("rows") or []:
            ws.append([_coerce(v) for v in list(row)[: max(len(cols), len(row))]])
        if sh.get("total_row") and cols and ws.max_row > 1:
            totals = ["TOTAL"]
            for ci in range(2, len(cols) + 1):
                col = get_column_letter(ci)
                numeric = any(isinstance(ws.cell(r, ci).value, (int, float)) for r in range(2, ws.max_row + 1))
                totals.append(f"=SUM({col}2:{col}{ws.max_row})" if numeric else "")
            ws.append(totals)
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
        for i, c in enumerate(cols, start=1):
            width = max([len(c)] + [len(str(ws.cell(r, i).value or "")) for r in range(2, min(ws.max_row, 200) + 1)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(str(path))
